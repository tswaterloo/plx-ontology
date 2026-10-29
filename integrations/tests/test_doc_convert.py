"""Portal's MarkItDown path for Office/PDF originals (DOC-07), isolated fixtures.

Every fixture is generated here. The adapter tests inject a fake with
MarkItDown's API (``MarkItDown(enable_plugins=...)``, ``convert_local(path)``
returning an object with ``markdown``/``text_content``); the fallback tests
pin "not installed" and always run; the real-package tests skip cleanly when
``markitdown`` is not importable. No test opens a network connection: socket
connects are made to fail while converting.
"""

from __future__ import annotations

from datetime import datetime, timedelta, timezone
from email.message import EmailMessage
from hashlib import sha256
import builtins
import importlib
import importlib.machinery
import importlib.util
import io
import json
import re
from pathlib import Path
import socket
import struct
import subprocess
import sys
import types
import unittest
from unittest.mock import patch
from xml.sax.saxutils import escape
from zipfile import ZipFile

from integrations import doc_convert
from integrations import office_text
from integrations import document_understanding as du
from integrations.document_understanding import extract_document
from ontology.core import ValidationError
from integrations.tests.office_fixtures import (DOC_RELS, PURCHASE_ORDER_BODY, PURCHASE_ORDER_MD,
                                                W, make_word)
from integrations.tests.office_fixtures import content_types as _types
from integrations.tests.office_fixtures import package_rels as _rels
from integrations.tests.office_fixtures import paragraph as _p
from ontology.core import ValidationError
from operating import use_case_parts


URL = "https://tenant.sharepoint.com/sites/purchasing/original"
DOCX = doc_convert.DOCX
PPTX = doc_convert.PPTX
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
BACKEND_PATCH = "integrations.doc_convert.installed_backend"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _importable(name: str) -> bool:
    try:
        importlib.import_module(name)
    except (KeyboardInterrupt, SystemExit):
        raise
    except BaseException:
        return False
    return True


# --- synthetic originals (stdlib only) -------------------------------------

def make_docx(paragraphs: list[str]) -> bytes:
    body = "".join(f"<w:p><w:r><w:t xml:space='preserve'>{escape(p)}</w:t></w:r></w:p>"
                   for p in paragraphs)
    buffer = io.BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml",
            "<Types xmlns='http://schemas.openxmlformats.org/package/2006/content-types'>"
            "<Default Extension='rels' ContentType='application/vnd.openxmlformats-package.relationships+xml'/>"
            "<Default Extension='xml' ContentType='application/xml'/>"
            "<Override PartName='/word/document.xml' ContentType="
            "'application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml'/></Types>")
        archive.writestr("_rels/.rels",
            "<Relationships xmlns='http://schemas.openxmlformats.org/package/2006/relationships'>"
            "<Relationship Id='rId1' Type='http://schemas.openxmlformats.org/officeDocument/2006/"
            "relationships/officeDocument' Target='word/document.xml'/></Relationships>")
        archive.writestr("word/document.xml",
            f"<w:document xmlns:w='{W}'><w:body>{body}</w:body></w:document>")
    return buffer.getvalue()


def make_pptx(slides: list[list[str]]) -> bytes:
    """A deck of one text box per slide (``make_deck``): a PPTX the stdlib converter
    opens, as MarkItDown only runs on one it does."""
    return make_deck([(_shape(2, lines), {}) for lines in slides])


def make_xlsx() -> bytes:
    buffer = io.BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml",
            "<Types xmlns='http://schemas.openxmlformats.org/package/2006/content-types'/>")
        archive.writestr("xl/workbook.xml",
            "<workbook xmlns='http://schemas.openxmlformats.org/spreadsheetml/2006/main' "
            "xmlns:r='http://schemas.openxmlformats.org/officeDocument/2006/relationships'>"
            "<sheets><sheet name='Quote' sheetId='1' r:id='rId1'/></sheets></workbook>")
        archive.writestr("xl/_rels/workbook.xml.rels",
            "<Relationships xmlns='http://schemas.openxmlformats.org/package/2006/relationships'>"
            "<Relationship Id='rId1' Type='http://schemas.openxmlformats.org/officeDocument/2006/"
            "relationships/worksheet' Target='worksheets/sheet1.xml'/></Relationships>")
        archive.writestr("xl/worksheets/sheet1.xml",
            "<worksheet xmlns='http://schemas.openxmlformats.org/spreadsheetml/2006/main'><sheetData>"
            "<row r='1'><c r='A1' t='inlineStr'><is><t>Quote Number</t></is></c>"
            "<c r='B1' t='inlineStr'><is><t>Q-001</t></is></c></row>"
            "<row r='2'><c r='A2' t='inlineStr'><is><t>Total</t></is></c>"
            "<c r='B2'><f>C2*2</f><v>24</v></c></row></sheetData></worksheet>")
    return buffer.getvalue()


def make_pdf(pages: list[list[str]]) -> bytes:
    """A small multi-page PDF with a real Helvetica text layer, built by hand."""
    def literal(text: str) -> str:
        return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    count = len(pages)
    objects = ["<< /Type /Catalog /Pages 2 0 R >>",
               "<< /Type /Pages /Kids [" + " ".join(f"{3 + 2 * i} 0 R" for i in range(count)) +
               f"] /Count {count} >>"]
    font = 3 + 2 * count
    for index, lines in enumerate(pages):
        stream = "".join(f"BT /F1 12 Tf 72 {720 - 20 * row} Td ({literal(line)}) Tj ET\n"
                         for row, line in enumerate(lines))
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                       f"/Contents {4 + 2 * index} 0 R /Resources << /Font << /F1 {font} 0 R >> >> >>")
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}endstream")
    objects.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
            ).encode()
    return bytes(out)


# --- a fake with MarkItDown's API ------------------------------------------

class FakeResult:
    def __init__(self, markdown: str):
        self.markdown = markdown

    @property
    def text_content(self) -> str:
        return self.markdown


class FakeMarkItDown:
    """Mimics ``markitdown.MarkItDown``: docx paragraphs, pptx slide markers,
    pdf pages ending in form feeds (pdfminer's output shape)."""

    constructed: list[dict] = []
    calls: list[tuple[str, dict, bool, bytes]] = []
    fail = False

    def __init__(self, **kwargs):
        FakeMarkItDown.constructed.append(kwargs)

    def convert_local(self, path, **kwargs):
        local = Path(path)
        FakeMarkItDown.calls.append((str(path), kwargs, local.is_absolute(), local.read_bytes()))
        if FakeMarkItDown.fail:
            raise ValueError("fake conversion failure")
        content = local.read_bytes()
        if local.suffix == ".docx":
            with ZipFile(io.BytesIO(content)) as archive:
                from xml.etree import ElementTree as ET
                root = ET.fromstring(archive.read("word/document.xml"))
            paragraphs = ["".join(t.text or "" for t in p.iter(f"{{{W}}}t"))
                          for p in root.iter(f"{{{W}}}p")]
            # markdownify (MarkItDown 0.1.x DOCX) escapes "_" and "*" only.
            return FakeResult("\n\n".join(p.replace("_", "\\_").replace("*", "\\*")
                                            for p in paragraphs))
        if local.suffix == ".pptx":
            from xml.etree import ElementTree as ET
            parts = []
            with ZipFile(io.BytesIO(content)) as archive:
                names = sorted(n for n in archive.namelist() if n.startswith("ppt/slides/slide"))
                for number, name in enumerate(names, start=1):
                    root = ET.fromstring(archive.read(name))
                    text = "\n".join(t.text or "" for t in root.iter(f"{{{A}}}t"))
                    parts.append(f"\n\n<!-- Slide number: {number} -->\n{text}")
            return FakeResult("".join(parts))
        if local.suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(content))
            return FakeResult("".join((page.extract_text() or "") + "\n\n\f"
                                      for page in reader.pages))
        raise ValueError("unsupported")


FAKE = doc_convert.Backend(FakeMarkItDown, "0.1.6-fake")


def no_network():
    """Fail any socket connect for the duration of a conversion."""
    def refuse(*_args, **_kwargs):
        raise AssertionError("network access attempted during conversion")
    return patch.multiple(socket.socket, connect=refuse, connect_ex=refuse)


def pypdf_rejects(pages: int | None, failure: str = "pdf_parse_error"):
    """pypdf cannot read the PDF (``failure``) and poppler counts ``pages``: the case
    MarkItDown reads a PDF in (one pypdf reads keeps pypdf's reading)."""
    return patch.multiple("integrations.document_understanding",
                          _pypdf_open=lambda content, max_pages: (None, failure, (failure,)),
                          _poppler_page_count=lambda content: pages)


def pypdf_reads_no_text():
    """pypdf opens the PDF but reads no text on any page (a font it cannot map): the
    case MarkItDown reads a PDF pypdf opens in. No page's ``extract_text`` runs."""
    def empty(source, reader, max_text, first=None):
        count = len(reader.pages)
        read = count if first is None else min(first, count)
        return du._paged(source, "pypdf-text",
                         [du.Page(n, "", "pypdf-text") for n in range(1, read + 1)],
                         document=reader,
                         skipped_pages=(read + 1, count) if count > read else None)
    return patch("integrations.document_understanding._pypdf_pages", empty)


def poppler_reads(texts: list[str]):
    """poppler's text layer reads ``texts``, one per page, without running poppler."""
    return patch("integrations.document_understanding._pdftotext_pages",
                 lambda content, source, max_pages, first_pages: du._paged(
                     source, "pdftotext-text",
                     [du.Page(n, t, "pdftotext-text") for n, t in enumerate(texts, 1)]))


def stored_provenance_valid(attributes):
    """The stored conversion record matches the stored text and original hash
    (``doc_convert.valid_record``; formerly checked inside the deleted
    integrations/transaction_evidence.py importer)."""
    extraction = attributes["extraction"]
    return doc_convert.valid_record(extraction["conversion"], text=attributes.get("text", ""),
                                    source_sha256=extraction["source_sha256"])


class ConverterTestCase(unittest.TestCase):
    def setUp(self):
        FakeMarkItDown.constructed.clear()
        FakeMarkItDown.calls.clear()
        FakeMarkItDown.fail = False

    def extract(self, content, mime, *, backend=FAKE, **options):
        with patch(BACKEND_PATCH, return_value=backend), no_network():
            return extract_document(content, mime, original_url=URL, **options)

    def assert_portal_call(self):
        """Exactly Portal's options: MarkItDown(enable_plugins=False).convert_local(path)."""
        self.assertEqual(FakeMarkItDown.constructed, [{"enable_plugins": False}])
        self.assertEqual(len(FakeMarkItDown.calls), 1)
        path, kwargs, absolute, _ = FakeMarkItDown.calls[0]
        self.assertEqual(kwargs, {})
        self.assertTrue(absolute)
        self.assertFalse(Path(path).exists(), "private temp copy must be deleted")

    def assert_record(self, document, content, fmt):
        record = document.conversion
        self.assertEqual(record["label"], "machine-converted")
        self.assertEqual((record["converter"], record["converter_version"], record["format"]),
                         ("markitdown", "0.1.6-fake", fmt))
        self.assertEqual(record["source_sha256"], sha256(content).hexdigest())
        self.assertEqual(record["source_sha256"], document.source_sha256)
        self.assertEqual(record["output_sha256"], sha256(document.text.encode()).hexdigest())
        self.assertEqual(record["options"], {"constructor": {"enable_plugins": False},
            "call": "convert_local", "input": "private-local-file",
            "result": ["markdown", "text_content"], "postprocess": "none"})
        self.assertEqual(record["options_sha256"], sha256(json.dumps(record["options"],
            sort_keys=True, separators=(",", ":")).encode()).hexdigest())
        self.assertEqual(record["upstream"],
                         "plx-customer-portal:scripts/markitdown/convert.py@5724575")
        return record


class MarkItDownAdapterTest(ConverterTestCase):
    def test_docx_is_converted_with_portal_options_and_cited_provenance(self):
        content = make_docx(["Petra PO: PPO-200", "Supplier: Acme Packaging",
                             "Item: BTL-100 | Qty: 100,000 bottles | Unit Price: USD 0.12"])
        document = self.extract(content, DOCX, filename="po.docx")
        self.assert_portal_call()
        self.assertEqual(FakeMarkItDown.calls[0][3], content)  # the exact bytes
        self.assertEqual((document.status, document.method), ("extracted", "markitdown"))
        self.assertTrue(document.text.startswith("Petra PO: PPO-200"))
        record = self.assert_record(document, content, "docx")
        self.assertEqual(record["spans"], [])
        po = next(c for c in document.claims if c["field"] == "petra_po")
        self.assertEqual(document.text[po["span"]["start"]:po["span"]["end"]], "PPO-200")
        self.assertEqual(po["confidence"], "literal")
        self.assertEqual(document.line_items[0]["unit_price"], 0.12)
        attrs = document.evidence_attributes(URL, title="po.docx")
        self.assertEqual(attrs["extraction"]["conversion"], record)
        self.assertEqual(attrs["extraction"]["method"], "markitdown")
        self.assertTrue(stored_provenance_valid(attrs))

    def test_an_abandoned_conversion_leaves_no_folder_of_customer_bytes(self):
        """A conversion its caller abandons (a budget; the eval watchdog before its
        ``os._exit``) has its private copy of the document removed at once."""
        import tempfile
        import threading
        from scripts import eval_oracle_dev as dev
        entered, release = threading.Event(), threading.Event()

        class Hanging(FakeMarkItDown):
            def convert_local(self, path, **kwargs):
                entered.set()
                release.wait(30)
                return super().convert_local(path, **kwargs)

        self.addCleanup(release.set)
        with tempfile.TemporaryDirectory() as tmp, patch.object(tempfile, "tempdir", tmp):
            self.assertIsNone(dev._within(0.5, doc_convert.convert, make_docx(["PPO-200"]),
                                          "docx", backend=doc_convert.Backend(Hanging, "x")))
            self.assertTrue(entered.is_set())
            self.assertEqual(list(Path(tmp).glob("plx-doc-convert-*")), [])

    def test_pptx_slide_spans_come_from_the_converter_markers(self):
        content = make_pptx([["Quote Number: Q-77"], ["Shipment reference: SR-9"]])
        document = self.extract(content, PPTX, filename="deck.pptx")
        self.assert_portal_call()
        record = self.assert_record(document, content, "pptx")
        self.assertEqual([(s["unit"], s["number"]) for s in record["spans"]],
                         [("slide", 1), ("slide", 2)])
        second = record["spans"][1]
        self.assertIn("SR-9", document.text[second["start"]:second["end"]])
        self.assertNotIn("Q-77", document.text[second["start"]:second["end"]])
        self.assertTrue(stored_provenance_valid(document.evidence_attributes(URL)))

    def test_octet_stream_office_files_are_recognised_by_name(self):
        document = self.extract(make_docx(["Petra PO: PPO-1"]), "application/octet-stream",
                                filename="PO.DOCX")
        self.assertEqual(document.method, "markitdown")
        self.assertIsNone(doc_convert.convertible_kind("application/octet-stream", "x.doc"))

    @unittest.skipUnless(_importable("pypdf"), "PDF page count needs pypdf")
    def test_pdf_page_spans_match_the_pdf_page_count(self):
        content = make_pdf([["Petra PO: PPO-200"], ["Shipment reference: SR-42"]])
        with pypdf_rejects(2):
            document = self.extract(content, "application/pdf", filename="po.pdf")
        self.assert_portal_call()
        self.assertEqual((document.status, document.method), ("extracted", "markitdown"))
        record = self.assert_record(document, content, "pdf")
        self.assertEqual([(s["unit"], s["number"]) for s in record["spans"]],
                         [("page", 1), ("page", 2)])
        self.assertEqual([p["page"] for p in document.pages], [1, 2])
        claim = next(c for c in document.claims if c["field"] == "shipment_reference")
        self.assertEqual(claim["page"], 2)
        self.assertEqual(document.warnings, ())
        self.assertTrue(stored_provenance_valid(document.evidence_attributes(URL)))

    @unittest.skipUnless(_importable("pypdf"), "PDF page count needs pypdf")
    def test_pdf_page_spans_are_never_guessed_the_pages_fall_back_and_say_so(self):
        class OnePiece(FakeMarkItDown):
            def convert_local(self, path, **kwargs):
                super().convert_local(path, **kwargs)
                return FakeResult("Petra PO: PPO-200 Shipment reference: SR-42")
        content = make_pdf([["Petra PO: PPO-200"], ["Shipment reference: SR-42"]])
        # MarkItDown gave no page boundaries (pypdf panicked; poppler read the text
        # layer): the next reader's pages, recorded as such.
        with pypdf_rejects(2, "parser_panicked"), \
                poppler_reads(["Petra PO: PPO-200", "Shipment reference: SR-42"]):
            document = self.extract(content, "application/pdf",
                                    backend=doc_convert.Backend(OnePiece, "0.1.6-fake"))
        self.assertEqual((document.status, document.method, document.conversion),
                         ("extracted", "pdftotext-text", None))
        self.assertEqual([p["page"] for p in document.pages], [1, 2])
        self.assertEqual(document.warnings, (du.PAGE_FALLBACK,))
        claim = next(c for c in document.claims if c["field"] == "shipment_reference")
        self.assertEqual(claim["page"], 2)
        # No reader gives pages (pypdf rejects the file; poppler counts its pages):
        # MarkItDown's text, unpaged.
        with patch("integrations.document_understanding._pypdf_open",
                   return_value=(None, "pdf_parse_error", ("pdf_parse_error:PdfReadError",))), \
                patch("integrations.document_understanding._poppler_page_count",
                      return_value=2):
            unpaged = self.extract(content, "application/pdf",
                                   backend=doc_convert.Backend(OnePiece, "0.1.6-fake"))
        self.assertEqual((unpaged.status, unpaged.method, unpaged.pages),
                         ("extracted", "markitdown", ()))
        self.assertEqual(unpaged.conversion["spans"], [])
        self.assertEqual(unpaged.warnings, ("converted_page_spans_unavailable",))
        self.assertTrue(all("page" not in claim for claim in unpaged.claims))

    @unittest.skipUnless(_importable("pypdf"), "PDF page count needs pypdf")
    def test_image_only_pdf_stays_on_the_ocr_route(self):
        document = self.extract(make_pdf([[]]), "application/pdf")
        self.assertEqual(document.status, "no_text")
        self.assertIn("image_only_or_empty_pdf_requires_ocr", document.warnings)
        self.assertEqual(document.conversion["label"], "machine-converted")

    @unittest.skipUnless(_importable("pypdf"), "encryption needs pypdf")
    def test_encrypted_pdf_is_refused_before_conversion(self):
        from pypdf import PdfReader, PdfWriter
        writer = PdfWriter()
        writer.append(PdfReader(io.BytesIO(make_pdf([["Petra PO: PPO-1"]]))))
        writer.encrypt("secret")
        buffer = io.BytesIO()
        writer.write(buffer)
        document = self.extract(buffer.getvalue(), "application/pdf")
        self.assertEqual((document.status, document.warnings), ("unsupported", ("encrypted_pdf",)))
        self.assertEqual(FakeMarkItDown.calls, [])
        self.assertIsNone(document.conversion)

    def test_pdf_over_the_page_bound_is_rejected_before_conversion(self):
        with patch("integrations.document_understanding._pypdf_open",
                   side_effect=ValidationError("PDF page count exceeds bound")):
            with self.assertRaises(ValidationError):
                self.extract(b"%PDF-1.4", "application/pdf")
        self.assertEqual(FakeMarkItDown.calls, [])

    def test_scanned_ocr_stays_workbench_work(self):
        stub = ("AO LOT #: AO-26005B", ({"page": 1, "start": 0, "end": 19},), "extracted",
                ("ocr_text_requires_identity_review",), frozenset({1}), 1)
        with patch("integrations.document_understanding._local_ocr_backend",
                   return_value="portable"), \
             patch("integrations.document_understanding._portable_pdf_text", return_value=stub):
            document = self.extract(b"%PDF-1.4", "application/pdf", local_pdf_ocr=True)
        self.assertEqual(FakeMarkItDown.calls, [])
        self.assertEqual((document.method, document.conversion), ("poppler-tesseract", None))
        self.assertEqual(document.ocr_pages, (1,))

    def test_xlsx_keeps_the_cell_transcription(self):
        document = self.extract(make_xlsx(), XLSX, filename="quote.xlsx")
        self.assertEqual(FakeMarkItDown.calls, [])
        self.assertEqual((document.method, document.conversion),
                         ("ooxml-cell-transcription", None))
        self.assertIn("B2=FORMULA(C2*2); cached=24", document.text)
        self.assertIn("xlsx_formulas_not_recalculated", document.warnings)

    def test_converter_errors_are_explicit_and_pdf_falls_back_to_its_text_layer(self):
        FakeMarkItDown.fail = True
        office = self.extract(make_docx(["Petra PO: PPO-1"]), DOCX)
        # The stdlib converter reads what MarkItDown could not; its error is kept.
        self.assertEqual((office.status, office.text, office.method),
                         ("extracted", "Petra PO: PPO-1", "plx-office-stdlib"))
        self.assertEqual(office.warnings, ("markitdown_convert_error:ValueError",))
        # Bytes the stdlib converter refuses never reach MarkItDown.
        office = self.extract(b"not a zip", DOCX)
        self.assertEqual((office.status, office.text, office.conversion), ("unsupported", "", None))
        self.assertEqual(office.warnings, ("docx_invalid_zip",))

    @unittest.skipUnless(_importable("pypdf"), "the PDF text layer needs pypdf")
    def test_a_pdf_markitdown_cannot_read_falls_back_to_its_text_layer(self):
        FakeMarkItDown.fail = True
        with pypdf_rejects(1, "parser_panicked"), poppler_reads(["Petra PO: PPO-9"]):
            pdf = self.extract(make_pdf([["Petra PO: PPO-9"]]), "application/pdf")
        self.assertEqual((pdf.status, pdf.method, pdf.conversion),
                         ("extracted", "pdftotext-text", None))
        self.assertEqual(pdf.warnings, ("markitdown_convert_error:ValueError",
                                        "parser_panicked"))

    @unittest.skipUnless(_importable("pypdf"), "the PDF is read by pypdf")
    def test_a_pdf_pypdf_reads_keeps_its_table_rows_with_markitdown_installed(self):
        """pdfminer (MarkItDown's PDF reader) puts a blank line between lines spaced 1.5x
        their font or more, and a blank line ends a table: a CoA with 20 pt rows of 12 pt
        text lost every test line. A PDF pypdf reads keeps pypdf's reading."""
        from ontology import document_fields

        class Pdfminer(FakeMarkItDown):
            def convert_local(self, path, **kwargs):
                text = super().convert_local(path, **kwargs).markdown
                return FakeResult("\f".join("\n\n".join(page.split("\n"))
                                             for page in text.split("\f")))
        coa = ["CERTIFICATE OF ANALYSIS", "Certificate No: COA-24-0917",
               "Supplier: Ningbo Chemical Co", "Lot No: GL240917",
               "| Test | Specification | Result | Unit | Method |", "| --- | --- | --- | --- | --- |",
               "| Assay | 99.0 - 101.0 | 99.7 | % | USP <591> |",
               "| Water | max 0.5 | 0.0? | % | KF |"]
        content = make_pdf([coa])
        backends = [doc_convert.Backend(Pdfminer, "0.1.8-fake")]
        installed = doc_convert.installed_backend()
        if installed is not None:  # the real package too, where it is installed
            backends.append(installed)
        for backend in backends:
            with self.subTest(converter=backend.version):
                FakeMarkItDown.calls.clear()
                with patch(BACKEND_PATCH, return_value=backend), no_network():
                    fields = document_fields.extract_fields(content, "application/pdf",
                                                            filename="coa.pdf")
                self.assertEqual([item["cells"]["test"]["value"]
                                  for item in fields["line_items"]], ["Assay", "Water"])
                self.assertEqual(FakeMarkItDown.calls, [])

    def test_no_result_text_is_a_conversion_failure_as_in_portal(self):
        class Empty(FakeMarkItDown):
            def convert_local(self, path, **kwargs):
                class Nothing:
                    markdown = None
                    text_content = None
                return Nothing()
        document = self.extract(make_docx(["x"]), DOCX,
                                backend=doc_convert.Backend(Empty, "0.1.6-fake"))
        self.assertEqual(document.warnings, ("markitdown_returned_no_text",))

    def test_a_mail_attachment_is_converted_as_its_own_document(self):
        # The attachment is a container member: converted on its own bytes, never as a
        # section of the mail's text (direction finding 11e322b9).
        attachment = make_docx(["Vendor confirmation: VC-300", "Customer PO: CPO_42"])
        message = EmailMessage()
        message["Subject"] = "Confirmation"
        message["From"] = "supplier@example.test"
        message["To"] = "buyer@example.test"
        message.set_content("See attached.")
        message.add_attachment(attachment, maintype="application",
            subtype="vnd.openxmlformats-officedocument.wordprocessingml.document",
            filename="confirmation.docx")
        document = self.extract(message.as_bytes(), "message/rfc822")
        self.assertEqual((document.method, document.conversion), ("rfc822-mime", None))
        self.assertNotIn("VC-300", document.text)
        [member] = use_case_parts.members(message.as_bytes(), "confirmation.eml")["files"]
        self.assertEqual((member["name"], member["data"]), ("confirmation.docx", attachment))
        converted = self.extract(member["data"], member["hint"])
        record = converted.conversion
        self.assertEqual(record["source_sha256"], sha256(attachment).hexdigest())
        by_field = {c["field"]: c for c in converted.claims}
        self.assertEqual((by_field["vendor_confirmation"]["value"],
                          by_field["vendor_confirmation"]["converted"]), ("VC-300", True))
        claim = by_field["customer_po"]
        self.assertEqual((claim["value"], claim["confidence"], claim["converted"]),
                         ("CPO_42", "literal", True))
        self.assertEqual(converted.text[claim["span"]["start"]:claim["span"]["end"]], "CPO\\_42")
        attrs = converted.evidence_attributes(URL)
        self.assertTrue(stored_provenance_valid(attrs))
        # A record of converted sections inside other text is never written: refused.
        sections = {"label": "contains-machine-converted-sections", "sections": [
            {**record, "attachment": "confirmation.docx", "start": 0,
             "end": len(converted.text)}]}
        self.assertFalse(doc_convert.valid_record(sections, text=converted.text,
                                                  source_sha256=record["source_sha256"]))

    def test_stored_provenance_must_match_the_stored_text_and_original(self):
        document = self.extract(make_docx(["Petra PO: PPO-200"]), DOCX)
        good = document.evidence_attributes(URL)
        tampered_text = json.loads(json.dumps(good))
        tampered_text["text"] = good["text"] + " "
        wrong_original = json.loads(json.dumps(good))
        wrong_original["extraction"]["conversion"]["source_sha256"] = "1" * 64
        wrong_options = json.loads(json.dumps(good))
        wrong_options["extraction"]["conversion"]["options"]["constructor"]["enable_plugins"] = True
        unlabelled = json.loads(json.dumps(good))
        del unlabelled["extraction"]["conversion"]["label"]
        for attrs in (tampered_text, wrong_original, wrong_options, unlabelled):
            self.assertFalse(stored_provenance_valid(attrs))
        self.assertTrue(stored_provenance_valid(good))


class ConvertedClaimsTest(ConverterTestCase):
    """F1/F4: Markdown escapes never become wrong literal values."""

    def test_docx_escapes_are_read_through_an_exact_offset_map(self):
        content = make_docx(["Petra PO: PO_12345", "Item: SKU_88*A",
                             "Supplier: Acme_Bottles *Ltd*", "Quote Number: Q_7"])
        document = self.extract(content, DOCX)
        # The stored text is the converter's output, escapes included.
        self.assertIn("PO\\_12345", document.text)
        self.assertIn("Acme\\_Bottles \\*Ltd\\*", document.text)
        by_field = {c["field"]: c for c in document.claims}
        self.assertEqual(by_field["petra_po"]["value"], "PO_12345")
        self.assertEqual(by_field["quote_number"]["value"], "Q_7")
        self.assertEqual(by_field["supplier"]["value"], "Acme_Bottles *Ltd*")
        # A value cut short by "*" is no claim, never the prefix "SKU".
        self.assertNotIn("product_code", by_field)
        for claim in document.claims:
            self.assertIs(claim["converted"], True)
            self.assertEqual((claim["confidence"], claim["interpretation"]), ("literal", False))
            cited = document.text[claim["span"]["start"]:claim["span"]["end"]]
            self.assertEqual(cited.replace("\\_", "_").replace("\\*", "*"), claim["value"])
        self.assertEqual(document.text[by_field["petra_po"]["span"]["start"]:
                                       by_field["petra_po"]["span"]["end"]], "PO\\_12345")
        self.assertTrue(stored_provenance_valid(document.evidence_attributes(URL)))

    def test_docx_markup_makes_a_claim_review_grade(self):
        class Bold(FakeMarkItDown):
            def convert_local(self, path, **kwargs):
                super().convert_local(path, **kwargs)
                return FakeResult("Supplier: **Acme** Ltd\nPetra PO: PPO-1")
        document = self.extract(make_docx(["x"]), DOCX,
                                backend=doc_convert.Backend(Bold, "0.1.6-fake"))
        by_field = {c["field"]: c for c in document.claims}
        self.assertEqual((by_field["supplier"]["confidence"],
                          by_field["supplier"]["interpretation"]), ("review", True))
        self.assertEqual(by_field["petra_po"]["confidence"], "literal")

    def test_pptx_claims_cite_their_slide_and_backslashes_are_review(self):
        content = make_pptx([["Quote Number: Q-77"], ["Supplier: Acme\\_Co"]])
        document = self.extract(content, PPTX)
        by_field = {c["field"]: c for c in document.claims}
        self.assertEqual(by_field["quote_number"]["slide"], 1)
        self.assertEqual(by_field["quote_number"]["confidence"], "literal")
        # PPTX output escapes nothing, so "\\_" may be the source's own text.
        self.assertEqual(by_field["supplier"]["slide"], 2)
        self.assertEqual(by_field["supplier"]["value"], "Acme\\_Co")
        self.assertEqual(by_field["supplier"]["confidence"], "review")
        self.assertTrue(stored_provenance_valid(document.evidence_attributes(URL)))

    def test_plain_text_claims_carry_no_converted_marker(self):
        document = extract_document(b"Petra PO: PO_1\n", "text/plain", original_url=URL)
        self.assertNotIn("converted", document.claims[0])
        self.assertEqual(document.claims[0]["value"], "PO_1")

    @unittest.skipUnless(_importable("pypdf"), "PDF page count needs pypdf")
    def test_pdf_claims_are_marked_converted(self):
        with pypdf_rejects(1):
            document = self.extract(make_pdf([["Petra PO: PPO-200"]]), "application/pdf")
        self.assertTrue(document.claims)
        self.assertTrue(all(c["converted"] is True and c["page"] == 1 for c in document.claims))

    def test_pdf_with_an_unknown_page_count_is_never_converted_whole(self):
        """Without a page count (pypdf rejects the file and
        poppler cannot count it) MarkItDown is never given the PDF: it could be over the
        page bound. Poppler's count bounds it: over the bound it is refused unconverted."""
        class Text(FakeMarkItDown):
            def convert_local(self, path, **kwargs):
                FakeMarkItDown.calls.append((str(path), kwargs, True, b""))
                return FakeResult("Petra PO: PPO-1\n\n\f")
        backend = doc_convert.Backend(Text, "0.1.6-fake")
        with patch("integrations.document_understanding._pypdf_open",
                   return_value=(None, "pdf_parse_error", ("pdf_parse_error:PdfReadError",))):
            with patch("integrations.document_understanding._poppler_page_count",
                       return_value=None):
                document = self.extract(b"%PDF-1.4", "application/pdf", backend=backend)
            self.assertEqual((document.status, document.method, document.failure,
                              document.warnings, FakeMarkItDown.calls),
                             ("unsupported", "none", "pdf_parse_error",
                              ("pdf_page_count_unknown_not_converted",
                               "pdf_parse_error:PdfReadError"), []))
            with patch("integrations.document_understanding._poppler_page_count",
                       return_value=du.MAX_PAGES + 1), \
                    self.assertRaisesRegex(ValidationError, "^PDF page count exceeds bound$"):
                self.extract(b"%PDF-1.4", "application/pdf", backend=backend)
            self.assertEqual(FakeMarkItDown.calls, [])
            with patch("integrations.document_understanding._poppler_page_count",
                       return_value=1):  # counted and within the bound: converted, paged
                document = self.extract(b"%PDF-1.4", "application/pdf", backend=backend)
        self.assertEqual((document.method, [p["page"] for p in document.pages]),
                         ("markitdown", [1]))


class FallbackWithoutMarkItDownTest(ConverterTestCase):
    """Always runs: the package is treated as not installed."""


    def test_pdf_uses_the_existing_text_layer_reader(self):
        from types import SimpleNamespace

        stub = du._paged("0" * 64, "pypdf-text", [du.Page(1, "Petra PO: PPO-9", "pypdf-text")])
        # A one-page reader: the page count is read to bound the PDF (first pages).
        with patch("integrations.document_understanding._pypdf_open",
                   return_value=(SimpleNamespace(pages=[object()]), None, ())), \
             patch("integrations.document_understanding._pypdf_pages",
                   return_value=stub) as reader:
            document = self.extract(b"%PDF-1.4", "application/pdf", backend=None)
        reader.assert_called_once()
        self.assertEqual((document.method, document.conversion, document.warnings),
                         ("pypdf-text", None, ()))

    @unittest.skipUnless(_importable("pypdf"), "real PDF text layer needs pypdf")
    def test_generated_pdf_through_pypdf(self):
        document = self.extract(make_pdf([["Petra PO: PPO-200"], ["Quote Number: Q-5"]]),
                                "application/pdf", backend=None)
        self.assertEqual((document.status, document.method), ("extracted", "pypdf-text"))
        self.assertEqual([p["page"] for p in document.pages], [1, 2])
        self.assertEqual(next(c for c in document.claims if c["field"] == "quote_number")["page"], 2)

    def test_xlsx_is_unchanged(self):
        document = self.extract(make_xlsx(), XLSX, backend=None)
        self.assertEqual((document.status, document.method), ("extracted", "ooxml-cell-transcription"))

    def test_broken_or_missing_package_means_unavailable(self):
        import builtins
        real_import = builtins.__import__

        class PanicException(BaseException):
            pass

        def fake_import(name, *args, **kwargs):
            if name == "markitdown":
                raise PanicException("native dependency panicked")
            return real_import(name, *args, **kwargs)

        doc_convert._import_backend.cache_clear()
        try:
            with patch("builtins.__import__", side_effect=fake_import):
                self.assertIsNone(doc_convert.installed_backend())
                self.assertFalse(doc_convert.available())
        finally:
            doc_convert._import_backend.cache_clear()


@unittest.skipUnless(_importable("pypdf"), "PDF page count needs pypdf")
class OperatorInboxOriginalTest(unittest.TestCase):
    def test_inbox_pdf_without_ocr_keeps_the_conversion_record(self):
        import tempfile
        from ontology.core import Actor
        from operating.case_originals import OriginalStore
        FakeMarkItDown.constructed.clear()
        FakeMarkItDown.calls.clear()
        FakeMarkItDown.fail = False
        operator = Actor("plx", frozenset({"staff:operations"}), "restricted")
        content = make_pdf([["Petra PO: PPO-200"], ["Quote Number: Q-5"]])
        with tempfile.TemporaryDirectory() as root:
            store = OriginalStore(root, local_ocr=False)
            inbox = store.inbox / "staff:operations"
            inbox.mkdir()
            (inbox / "po.pdf").write_bytes(content)
            with patch(BACKEND_PATCH, return_value=FAKE), no_network(), pypdf_rejects(2):
                attached = store.attach(operator, "po.pdf")
            extraction = attached["extraction"]
            self.assertEqual(extraction["method"], "markitdown")
            self.assertEqual(extraction["conversion"]["label"], "machine-converted")
            self.assertEqual(extraction["conversion"]["source_sha256"], attached["sha256"])
            self.assertEqual(len(attached["page_citations"]), 2)
            page2 = store.resolve(attached["page_citations"][1], operator)
            self.assertIn("Q-5", page2["text"])
            self.assertEqual(page2["text_source"], "machine_converted")
            self.assertEqual(page2["extraction"]["conversion"]["converter"], "markitdown")
            with patch(BACKEND_PATCH, return_value=FAKE), no_network(), pypdf_rejects(2):
                same = store.reextract(operator, attached["citation"], "repeat")
            self.assertTrue(same["duplicate"])
            newer = doc_convert.Backend(FakeMarkItDown, "0.1.9-fake")
            with patch(BACKEND_PATCH, return_value=newer), no_network(), pypdf_rejects(2):
                upgraded = store.reextract(operator, attached["citation"], "converter upgraded")
            self.assertFalse(upgraded["duplicate"])


@unittest.skipUnless(doc_convert.available(), "markitdown is not installed (optional)")
class RealMarkItDownTest(unittest.TestCase):
    """Runs only where the optional package is installed, as Portal pins it."""

    def setUp(self):
        doc_convert._import_backend.cache_clear()
        self.backend = doc_convert.installed_backend()
        self.assertIsNotNone(self.backend)

    def tearDown(self):
        doc_convert._import_backend.cache_clear()

    def check(self, content, mime, expected):
        with no_network():
            document = extract_document(content, mime, original_url=URL)
        from importlib.metadata import version
        self.assertEqual(document.method, "markitdown")
        self.assertEqual(document.conversion["converter_version"], version("markitdown"))
        self.assertEqual(document.conversion["source_sha256"], sha256(content).hexdigest())
        self.assertIn(expected, document.text)
        self.assertTrue(stored_provenance_valid(document.evidence_attributes(URL)))
        return document

    def test_docx(self):
        self.require_docx()
        if _importable("docx"):
            import docx
            built = docx.Document()
            built.add_paragraph("Petra PO: PPO-200")
            buffer = io.BytesIO()
            built.save(buffer)
            content = buffer.getvalue()
        else:
            content = make_docx(["Petra PO: PPO-200"])
        document = self.check(content, DOCX, "PPO-200")
        self.assertTrue(any(c["value"] == "PPO-200" for c in document.claims))

    def require_docx(self):
        if not doc_convert.available("docx"):
            self.skipTest("markitdown's docx extra is missing: "
                          + "+".join(doc_convert.missing_extras("docx")))

    def test_docx_underscore_and_asterisk_identifiers(self):
        self.require_docx()
        if _importable("docx"):
            import docx
            built = docx.Document()
            for line in ("Petra PO: PO_12345", "Item: SKU_88*A", "Supplier: Acme_Bottles *Ltd*"):
                built.add_paragraph(line)
            buffer = io.BytesIO()
            built.save(buffer)
            content = buffer.getvalue()
        else:
            content = make_docx(["Petra PO: PO_12345", "Item: SKU_88*A",
                                 "Supplier: Acme_Bottles *Ltd*"])
        document = self.check(content, DOCX, "PO\\_12345")
        by_field = {c["field"]: c for c in document.claims}
        self.assertEqual(by_field["petra_po"]["value"], "PO_12345")
        self.assertEqual(by_field["supplier"]["value"], "Acme_Bottles *Ltd*")
        self.assertNotIn("product_code", by_field)
        self.assertTrue(all(c["converted"] is True for c in document.claims))

    @unittest.skipUnless(_importable("pptx"), "python-pptx builds the real deck")
    def test_pptx(self):
        import pptx
        deck = pptx.Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = "Quote Number: Q-77"
        buffer = io.BytesIO()
        deck.save(buffer)
        document = self.check(buffer.getvalue(), PPTX, "Q-77")
        self.assertEqual(document.conversion["spans"][0]["unit"], "slide")
        self.assertEqual(next(c for c in document.claims if c["field"] == "quote_number")["slide"], 1)

    def test_pdf(self):
        with pypdf_rejects(2):
            self.check(make_pdf([["Petra PO: PPO-200"], ["Quote Number: Q-5"]]),
                       "application/pdf", "Q-5")

    @unittest.skipUnless(_importable("pptx"), "python-pptx reads the deck")
    def test_stdlib_converter_writes_markitdown_output(self):
        """The stdlib converter's output is MarkItDown's for the same original."""
        import tempfile
        for kind, content in (("docx", make_word(PURCHASE_ORDER_BODY)),
                              ("pptx", make_deck(QUOTE_DECK))):
            if not doc_convert.available(kind):
                with self.subTest(kind=kind):
                    self.skipTest(f"markitdown {kind} extra missing: "
                                  + "+".join(doc_convert.missing_extras(kind)))
                continue
            with tempfile.TemporaryDirectory() as directory, no_network():
                path = Path(directory) / f"original.{kind}"
                path.write_bytes(content)
                converter = self.backend.factory(enable_plugins=False)
                expected = converter.convert_local(str(path)).markdown
            self.assertEqual(doc_convert.convert_stdlib(content, kind).text, expected, kind)


class PerFormatReadinessTest(unittest.TestCase):
    """markitdown imports but its docx extra (mammoth, lxml) does not: docx is not ready."""

    @staticmethod
    def without(*names):
        """Only ``names`` are missing; every other converter extra reads as present, so the
        outcome does not depend on which extras (lxml, pdfminer, ...) this host has."""
        real = importlib.util.find_spec

        def find_spec(name, *a, **k):
            if name in names:
                return None
            if any(name in extras for extras in doc_convert._REQUIRES.values()):
                return importlib.machinery.ModuleSpec(name, None)
            return real(name, *a, **k)
        return patch("importlib.util.find_spec", find_spec)

    def setUp(self):
        class Double:  # stands for the imported package; never reached for docx here
            def __init__(self, **_):
                raise AssertionError("markitdown must not run for a format that is not ready")
        self.backend = doc_convert.Backend(Double, "0.1.6", real=True)
        self.content = make_word(PURCHASE_ORDER_BODY)

    def test_docx_is_not_ready_and_convert_falls_back_with_a_reason(self):
        with patch.object(doc_convert, "installed_backend", return_value=self.backend), \
                self.without("mammoth"):
            self.assertTrue(doc_convert.available())
            self.assertFalse(doc_convert.available("docx"))
            self.assertEqual(doc_convert.missing_extras("docx"), ("mammoth",))
            converted = doc_convert.convert(self.content, "docx", backend=self.backend)
            with patch.object(doc_convert, "convert_stdlib",
                              wraps=doc_convert.convert_stdlib) as stdlib:
                reading = du._office_reading(self.content, "docx")
            self.assertEqual(stdlib.call_count, 1)  # converted once, not again on fallback
        self.assertEqual(converted.converter, doc_convert.STDLIB_CONVERTER)
        self.assertEqual(converted.text, doc_convert.convert_stdlib(self.content, "docx").text)
        self.assertIn("markitdown_docx_not_ready:missing_mammoth", converted.warnings)
        self.assertEqual(reading.method, doc_convert.STDLIB_CONVERTER)
        self.assertIn("markitdown_docx_not_ready:missing_mammoth", reading.warnings)

    def test_pdf_without_pdfminer_fails_closed_with_the_reason(self):
        with self.without("pdfminer"), self.assertRaises(doc_convert.ConversionFailed) as raised:
            doc_convert.convert(b"%PDF-1.4", "pdf", backend=self.backend)
        self.assertEqual(str(raised.exception), "markitdown_pdf_not_ready:missing_pdfminer")

    def test_a_ready_format_still_uses_markitdown(self):
        class Ok:
            def __init__(self, **_): pass
            def convert_local(self, path):
                return type("R", (), {"markdown": "x"})()
        backend = doc_convert.Backend(Ok, "0.1.6", real=True)
        with self.without("nothing_else"), patch.object(doc_convert, "_REQUIRES", {}):
            self.assertEqual(doc_convert.convert(self.content, "docx", backend=backend).converter,
                             doc_convert.CONVERTER)

    def test_real_backend_docx_tests_and_egress_probe_skip_when_docx_is_not_ready(self):
        import unittest as u
        from tests import test_doc_backend_egress as egress
        for case in (RealMarkItDownTest("test_docx"),
                     egress.DocumentBackendEgressTest("test_a_real_conversion_opens_no_connection")):
            result = u.TestResult()
            with patch.object(doc_convert, "installed_backend", return_value=self.backend), \
                    self.without("mammoth"):
                case.run(result)
            self.assertEqual((len(result.skipped), result.errors, result.failures),
                             (1, [], []), case.id())


# --- full OOXML/CFB originals for the stdlib converter (stdlib only) --------

PML = "http://schemas.openxmlformats.org/presentationml/2006/main"
CT_PML = "application/vnd.openxmlformats-officedocument.presentationml"


def _utf16_doctype_word() -> bytes:
    """A DOCX whose document part declares an entity in UTF-16, where a byte search
    for ``<!DOCTYPE`` finds nothing."""
    buffer = io.BytesIO()
    with ZipFile(io.BytesIO(make_word(_p("x")))) as source, ZipFile(buffer, "w") as archive:
        for name in source.namelist():
            data = source.read(name)
            if name == "word/document.xml":
                data = ("<?xml version='1.0' encoding='UTF-16'?><!DOCTYPE w:document "
                        "[<!ENTITY po 'PPO-1'>]>" + data.decode().replace("x<", "&po;<")
                        ).encode("utf-16")
            archive.writestr(name, data)
    return buffer.getvalue()


def _shape(sid: int, text: list[str], *, ph: str = "", y: int | None = None) -> str:
    nvpr = f"<p:nvPr><p:ph type='{ph}'/></p:nvPr>" if ph else "<p:nvPr/>"
    xfrm = (f"<a:xfrm><a:off x='0' y='{y}'/><a:ext cx='100' cy='100'/></a:xfrm>"
            if y is not None else "")
    paragraphs = "".join(f"<a:p><a:r><a:t>{escape(line)}</a:t></a:r></a:p>" for line in text)
    box = "" if ph else " txBox='1'"
    return (f"<p:sp><p:nvSpPr><p:cNvPr id='{sid}' name='Shape {sid}'/><p:cNvSpPr{box}/>{nvpr}"
            f"</p:nvSpPr><p:spPr>{xfrm}</p:spPr><p:txBody><a:bodyPr/>{paragraphs}</p:txBody></p:sp>")


def _deck_table(sid: int, rows: list[list[str]], y: int) -> str:
    cells = lambda row: "".join(  # noqa: E731
        f"<a:tc><a:txBody><a:bodyPr/><a:p><a:r><a:t>{escape(c)}</a:t></a:r></a:p></a:txBody></a:tc>"
        for c in row)
    return (f"<p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id='{sid}' name='Table {sid}'/>"
            "<p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>"
            f"<p:xfrm><a:off x='0' y='{y}'/><a:ext cx='100' cy='100'/></p:xfrm><a:graphic>"
            "<a:graphicData uri='http://schemas.openxmlformats.org/drawingml/2006/table'><a:tbl>"
            "<a:tblGrid>" + "<a:gridCol w='100'/>" * len(rows[0]) + "</a:tblGrid>" +
            "".join(f"<a:tr h='10'>{cells(row)}</a:tr>" for row in rows) +
            "</a:tbl></a:graphicData></a:graphic></p:graphicFrame>")


def make_deck(slides: list[tuple[str, dict]]) -> bytes:
    """A PPTX python-pptx opens: one master and layout (the title placeholder's
    position lives on the master), slides of shapes, optional speaker notes."""
    ns = f"xmlns:a='{A}' xmlns:p='{PML}' xmlns:r='{DOC_RELS}'"
    tree = lambda shapes: (  # noqa: E731
        "<p:cSld><p:spTree><p:nvGrpSpPr><p:cNvPr id='1' name=''/><p:cNvGrpSpPr/><p:nvPr/>"
        f"</p:nvGrpSpPr><p:grpSpPr/>{shapes}</p:spTree></p:cSld>")
    types = {"ppt/presentation.xml": CT_PML + ".presentation.main+xml",
             "ppt/slideMasters/slideMaster1.xml": CT_PML + ".slideMaster+xml",
             "ppt/slideLayouts/slideLayout1.xml": CT_PML + ".slideLayout+xml"}
    buffer = io.BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr("_rels/.rels", _rels(("rId1", "officeDocument", "ppt/presentation.xml")))
        archive.writestr("ppt/slideMasters/slideMaster1.xml",
            f"<p:sldMaster {ns}>" + tree(_shape(2, [], ph="title", y=500)) +
            "<p:sldLayoutIdLst><p:sldLayoutId id='2147483649' r:id='rId1'/></p:sldLayoutIdLst>"
            "</p:sldMaster>")
        archive.writestr("ppt/slideMasters/_rels/slideMaster1.xml.rels",
                         _rels(("rId1", "slideLayout", "../slideLayouts/slideLayout1.xml")))
        archive.writestr("ppt/slideLayouts/slideLayout1.xml",
                         f"<p:sldLayout {ns}>" + tree(_shape(2, [], ph="title")) + "</p:sldLayout>")
        archive.writestr("ppt/slideLayouts/_rels/slideLayout1.xml.rels",
                         _rels(("rId1", "slideMaster", "../slideMasters/slideMaster1.xml")))
        listed, rels = "", [("rId1", "slideMaster", "slideMasters/slideMaster1.xml")]
        for number, (shapes, extra) in enumerate(slides, start=1):
            name = f"ppt/slides/slide{number}.xml"
            types[name] = CT_PML + ".slide+xml"
            listed += f"<p:sldId id='{255 + number}' r:id='rId{number + 1}'/>"
            rels.append((f"rId{number + 1}", "slide", f"slides/slide{number}.xml"))
            archive.writestr(name, f"<p:sld {ns}>" + tree(shapes) + "</p:sld>")
            slide_rels = [("rId1", "slideLayout", "../slideLayouts/slideLayout1.xml")]
            if extra.get("notes"):
                notes = f"ppt/notesSlides/notesSlide{number}.xml"
                types[notes] = CT_PML + ".notesSlide+xml"
                slide_rels.append(("rId2", "notesSlide", f"../notesSlides/notesSlide{number}.xml"))
                archive.writestr(notes, f"<p:notes {ns}>" +
                                 tree(_shape(2, [extra["notes"]], ph="body")) + "</p:notes>")
                archive.writestr(f"ppt/notesSlides/_rels/notesSlide{number}.xml.rels",
                                 _rels(("rId1", "slide", f"../slides/slide{number}.xml")))
            archive.writestr(f"ppt/slides/_rels/slide{number}.xml.rels", _rels(*slide_rels))
        archive.writestr("ppt/presentation.xml",
            f"<p:presentation {ns}><p:sldMasterIdLst><p:sldMasterId id='2147483648' r:id='rId1'/>"
            f"</p:sldMasterIdLst><p:sldIdLst>{listed}</p:sldIdLst><p:sldSz cx='9144000' "
            "cy='6858000'/><p:notesSz cx='6858000' cy='9144000'/></p:presentation>")
        archive.writestr("ppt/_rels/presentation.xml.rels", _rels(*rels))
        archive.writestr("[Content_Types].xml", _types(types))
    return buffer.getvalue()


QUOTE_DECK = [
    (_shape(2, ["Quote Number: Q-77"], y=2000) + _shape(3, ["Quote Q_77"], ph="title"), {}),
    (_deck_table(4, [["SKU", "Qty"], ["BTL_1", "5"]], 3000) + _shape(3, ["Lines"], ph="title") +
     _shape(5, ["Draft"], y=100), {"notes": "Petra PO: PPO-9"}),
]
QUOTE_DECK_MD = ("<!-- Slide number: 1 -->\n# Quote Q_77\nQuote Number: Q-77\n\n"
                 "<!-- Slide number: 2 -->\nDraft\n# Lines\n| SKU | Qty |\n| --- | --- |\n"
                 "| BTL\\_1 | 5 |\n\n### Notes:\nPetra PO: PPO-9")


def make_cfb(tree: dict) -> bytes:
    """A version-3 Compound File ([MS-CFB]): ``{name: bytes | {...}}``, streams under
    4096 bytes in the mini stream, larger ones in regular sectors."""
    entries: list[dict] = []

    def add(name: str, node) -> int:
        index = len(entries)
        entries.append({})
        kids = [add(key, value) for key, value in node.items()] if isinstance(node, dict) else []
        entries[index] = {"name": name, "kind": 1 if isinstance(node, dict) else 2,
                          "data": b"" if isinstance(node, dict) else node, "kids": kids}
        return index

    end, free = 0xFFFFFFFE, 0xFFFFFFFF
    add("Root Entry", tree)
    entries[0]["kind"] = 5
    mini, mini_fat, big = bytearray(), [], []
    for entry in entries:
        data = entry["data"]
        if entry["kind"] == 2 and len(data) < 4096:
            count = -(-len(data) // 64)
            entry["start"] = len(mini_fat) if count else end
            mini_fat += [len(mini_fat) + i + 1 for i in range(count - 1)] + ([end] if count else [])
            mini += data + b"\0" * (-len(data) % 64)
        elif entry["kind"] == 2:
            big.append(entry)
    regions = [("dir", -(-len(entries) * 128 // 512)), ("minifat", -(-len(mini_fat) * 4 // 512)),
               ("mini", -(-len(mini) // 512))] + [(id(e), -(-len(e["data"]) // 512)) for e in big]
    other = sum(count for _, count in regions)
    fat_count = 1
    while fat_count * 128 < fat_count + other:
        fat_count += 1
    fat, starts, cursor = [0xFFFFFFFD] * fat_count, {}, fat_count
    for key, count in regions:
        starts[key] = cursor if count else end
        fat += [cursor + i + 1 for i in range(count - 1)] + ([end] if count else [])
        cursor += count
    fat += [free] * (fat_count * 128 - len(fat))
    for entry in big:
        entry["start"] = starts[id(entry)]
    entries[0]["start"], entries[0]["data"] = starts["mini"], bytes(mini)
    for entry in entries:
        entry.setdefault("left", free), entry.setdefault("right", free)
        entry["child"] = entry["kids"][0] if entry["kids"] else free
        for first, second in zip(entry["kids"], entry["kids"][1:]):
            entries[first]["right"] = second
    directory = b""
    for entry in entries:
        name = entry["name"].encode("utf-16-le") + b"\0\0"
        directory += (name.ljust(64, b"\0") + struct.pack("<HBB3I", len(name), entry["kind"], 1,
                      entry["left"], entry["right"], entry["child"]) + b"\0" * 36 +
                      struct.pack("<IQ", entry.get("start", end), len(entry["data"])))
    header = (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 16 +
              struct.pack("<HHHHH6xIIIIIIIII", 0x3E, 3, 0xFFFE, 9, 6, 0, fat_count, starts["dir"],
                          0, 4096, starts["minifat"], -(-len(mini_fat) * 4 // 512), end, 0) +
              struct.pack("<109I", *(list(range(fat_count)) + [free] * (109 - fat_count))))
    pad = lambda raw: raw + b"\0" * (-len(raw) % 512)  # noqa: E731
    body = (struct.pack(f"<{len(fat)}I", *fat) + pad(directory) +
            pad(struct.pack(f"<{len(mini_fat)}I", *mini_fat)) + pad(bytes(mini)) +
            b"".join(pad(entry["data"]) for entry in big))
    return header + body


def _unicode(text: str) -> bytes:
    return text.encode("utf-16-le")


def msg_tree(attachments: list[tuple[str, str, bytes | dict]], rtf: bytes | None = None, *,
             body: str = "Petra PO: PPO-200\r\nSee the attached quote.", header: int = 32
             ) -> dict:
    """An Outlook message's storage tree: subject, sender, recipients, submit time, plain
    body (or only the compressed RTF body ``rtf``) and attachments ``(filename, mime type,
    content)``: by value (bytes), or an embedded message (a ``msg_tree(..., header=24)``)."""
    sent = (datetime(2026, 9, 30, 10, 0, tzinfo=timezone.utc) -
            datetime(1601, 1, 1, tzinfo=timezone.utc)) // timedelta(microseconds=1) * 10
    tree = {"__properties_version1.0": b"\0" * header + struct.pack("<IIQ", 0x00390040, 6, sent),
            "__substg1.0_0037001F": _unicode("PO PPO-200 confirmation"),
            "__substg1.0_0C1A001F": _unicode("Acme Sales"),
            "__substg1.0_0C1F001F": _unicode("sales@acme.example"),
            "__substg1.0_0E04001F": _unicode("Purchasing"),
            "__substg1.0_1000001F": _unicode(body),
            "__recip_version1.0_#00000000": {"__properties_version1.0": b"\0" * 8}}
    if rtf is not None:
        del tree["__substg1.0_1000001F"]
        tree["__substg1.0_10090102"] = rtf
    for index, (name, mime, data) in enumerate(attachments):
        embedded = isinstance(data, dict)   # attach method 5 (embedded message), else 1
        tree[f"__attach_version1.0_#{index:08X}"] = {
            "__properties_version1.0": b"\0" * 8 + struct.pack("<IIQ", 0x37050003, 6,
                                                               5 if embedded else 1),
            "__substg1.0_3707001F": _unicode(name), "__substg1.0_370E001F": _unicode(mime),
            "__substg1.0_3701000D" if embedded else "__substg1.0_37010102": data}
    return tree


def make_msg(attachments: list[tuple[str, str, bytes | dict]], rtf: bytes | None = None,
             **fields) -> bytes:
    """An Outlook .msg (:func:`msg_tree`)."""
    return make_cfb(msg_tree(attachments, rtf, **fields))


class StdlibConverterTest(ConverterTestCase):
    """Without MarkItDown, DOCX/PPTX convert in its shape; failures are explicit."""

    def test_docx_converts_in_markitdown_shape_with_cited_provenance(self):
        content = make_word(PURCHASE_ORDER_BODY)
        document = self.extract(content, DOCX, backend=None)
        self.assertEqual((document.status, document.method, document.text),
                         ("extracted", "plx-office-stdlib", PURCHASE_ORDER_MD))
        record = document.conversion
        self.assertEqual((record["converter"], record["converter_version"], record["format"],
                          record["source_sha256"], record["spans"]),
                         ("plx-office-stdlib", doc_convert.STDLIB_VERSION, "docx",
                          sha256(content).hexdigest(), []))
        self.assertEqual(record["options"], doc_convert.STDLIB_OPTIONS)
        claim = next(c for c in document.claims if c["field"] == "petra_po")
        self.assertEqual((claim["value"], claim["converted"]), ("PO_12345", True))
        self.assertEqual(document.text[claim["span"]["start"]:claim["span"]["end"]], "PO\\_12345")
        attributes = document.evidence_attributes(URL)
        self.assertTrue(stored_provenance_valid(attributes))
        # A record claiming MarkItDown's options for the stdlib output is refused.
        attributes["extraction"]["conversion"]["options"] = doc_convert.OPTIONS
        self.assertFalse(stored_provenance_valid(attributes))
        boxed = self.extract(make_word(_p("Petra PO: PPO-1", extra=(
            f"<w:drawing><w:txbxContent>{_p('hidden')}</w:txbxContent></w:drawing>"))), DOCX,
            backend=None)
        self.assertEqual((boxed.text, boxed.warnings),
                         ("Petra PO: PPO-1", ("docx_text_boxes_not_converted",)))

    def test_pptx_converts_in_markitdown_shape_and_claims_cite_their_slide(self):
        document = self.extract(make_deck(QUOTE_DECK), PPTX, backend=None)
        self.assertEqual((document.status, document.method, document.text),
                         ("extracted", "plx-office-stdlib", QUOTE_DECK_MD))
        spans = document.conversion["spans"]
        self.assertEqual([(s["unit"], s["number"]) for s in spans], [("slide", 1), ("slide", 2)])
        self.assertEqual(spans[1]["end"], len(document.text))
        slides = {c["field"]: c["slide"] for c in document.claims}
        self.assertEqual(slides, {"quote_number": 1, "petra_po": 2})
        self.assertTrue(stored_provenance_valid(document.evidence_attributes(URL)))

    def test_unreadable_files_fail_closed_with_an_explicit_status(self):
        zip_without_document = io.BytesIO()
        with ZipFile(zip_without_document, "w") as archive:
            archive.writestr("[Content_Types].xml", _types({}))
        encrypted = make_cfb({"EncryptionInfo": b"\x04\0" * 8, "EncryptedPackage": b"\0" * 64})
        cases = [(b"not a zip", DOCX, "docx_invalid_zip"),
                 (zip_without_document.getvalue(), PPTX, "pptx_missing_part"),
                 (encrypted, DOCX, "docx_encrypted"),
                 (encrypted, XLSX, "xlsx_encrypted"),
                 (make_word("<w:p>"), DOCX, "docx_invalid_xml"),
                 (_utf16_doctype_word(), DOCX, "docx_dtd_not_permitted"),
                 (make_word(_p("x")).replace(b"<w:body>", b"<w:body><w:p>"), DOCX,
                  "docx_invalid_zip"),  # a member whose bytes no longer match its header
                 (make_xlsx().replace(b"Quote Number", b"Quote Numbex"), XLSX, "xlsx_invalid_zip"),
                 (b"\xd0\xcf\x11\xe0 truncated", "application/vnd.ms-outlook",
                  "msg_not_compound_file"),
                 (encrypted, "application/vnd.ms-outlook", "msg_not_outlook_message"),
                 (b"PK\x03\x04", "application/zip", "unsupported_mime_type:application/zip")]
        class AnyBytes(FakeMarkItDown):
            # MarkItDown's plain-text fallback reads bytes no Office converter opens.
            def convert_local(self, path, **kwargs):
                FakeMarkItDown.calls.append((str(path), kwargs, True, b""))
                return FakeResult(Path(path).read_bytes().decode("latin-1"))

        lenient = doc_convert.Backend(AnyBytes, "0.1.8-fake")
        for content, mime, reason in cases:
            office = mime in (DOCX, PPTX)
            for backend in (None, lenient) if office else (None,):
                with self.subTest(reason=reason, markitdown=backend is not None):
                    FakeMarkItDown.calls.clear()
                    document = self.extract(content, mime, backend=backend)
                    self.assertEqual((document.status, document.text, document.claims,
                                      document.warnings, document.conversion),
                                     ("unsupported", "", (), (reason,), None))
                    self.assertEqual(FakeMarkItDown.calls, [])

    def test_crafted_list_levels_and_style_chains_finish_fast(self):
        # A numbering level far past OOXML's 0..8 once drove an indent loop of that
        # many steps; a long basedOn chain was walked again for every paragraph.
        # The work is bounded by counting steps, not by the wall clock, so the bound
        # holds at any machine load: basedOn-chain steps are style lookups, and a list
        # item's per-level steps are its level.
        deep_level = make_word(
            _p("Top", bullet=True) + "<w:p><w:pPr><w:numPr><w:ilvl w:val='99999999999'/>"
            "<w:numId w:val='1'/></w:numPr></w:pPr><w:r><w:t>Deep</w:t></w:r></w:p>",
            levels="<w:lvl w:ilvl='99999999999'><w:numFmt w:val='decimal'/></w:lvl>")
        def chain(length: int, *, looped: bool = False) -> str:
            return "".join(f"<w:style w:type='paragraph' w:styleId='S{i}'><w:basedOn "
                           f"w:val='S{(i - 1) % length if looped else i - 1}'/></w:style>"
                           for i in range(length))

        long_chain = make_word("".join(_p(f"p{i}", style=f"S{i * 20}") for i in range(1500)),
                               styles=chain(30_000))
        long_cycle = make_word(_p("Looped", style="S0"), styles=chain(30_000, looped=True))
        cycle = make_word(_p("Looped", style="A"), styles=(
            "<w:style w:type='paragraph' w:styleId='A'><w:basedOn w:val='B'/></w:style>"
            "<w:style w:type='paragraph' w:styleId='B'><w:basedOn w:val='A'/><w:pPr>"
            "<w:numPr><w:numId w:val='1'/></w:numPr></w:pPr></w:style>"))
        steps = [0]
        levels: list[int] = []

        class CountingStyles(dict):
            def __getitem__(self, key):
                steps[0] += 1
                return super().__getitem__(key)

            def get(self, key, default=None):
                steps[0] += 1
                return super().get(key, default)

        reader_init, reader_list = office_text._Docx.__init__, office_text._Docx._list
        def counting_init(reader, package):
            reader_init(reader, package)
            reader.styles = CountingStyles(reader.styles)

        def bounded_list(reader, ppr, style_id):
            # Record each item's level; one past 0..8 is dropped before the renderer
            # can loop over it, so a regression fails the level bound, not a timeout.
            listed = reader_list(reader, ppr, style_id)
            if listed is not None:
                levels.append(listed[1])
            return listed if listed is None or listed[1] <= 8 else None

        def extract_counted(content: bytes):
            steps[0] = 0
            with patch.object(office_text._Docx, "__init__", counting_init), \
                    patch.object(office_text._Docx, "_list", bounded_list):
                document = self.extract(content, DOCX, backend=None)
            return document, steps[0]

        results = [extract_counted(content)
                   for content in (deep_level, long_chain, long_cycle, cycle)]
        self.assertEqual([(d.status, d.text, d.warnings) for d, _ in results],
                         [("extracted", "* Top\n\nDeep", ()),
                          ("unsupported", "", ("docx_exceeds_bound",)),
                          ("unsupported", "", ("docx_exceeds_bound",)),
                          ("extracted", "* Looped", ())])
        self.assertLessEqual(max(levels), 8)
        # A 30,000-style chain costs about the depth bound, not its length.
        for _, count in results[1:3]:
            self.assertLessEqual(count, 3 * office_text.MAX_STYLE_DEPTH)
        # n paragraphs each on its own style of one n-deep chain: each style is
        # resolved once, so the steps grow linearly in n (re-walking is quadratic).
        for n in (16, office_text.MAX_STYLE_DEPTH):
            with self.subTest(n=n):
                content = make_word("".join(_p(f"p{i}", style=f"S{i}") for i in range(n)),
                                    styles=chain(n))
                document, count = extract_counted(content)
                self.assertEqual(document.status, "extracted")
                # Every style is resolved at least once, however it is looked up.
                self.assertGreaterEqual(count, n)
                self.assertLessEqual(count, 4 * n)

    def test_outlook_msg_text_is_its_own_and_attachments_are_converted_apart(self):
        word = make_word(PURCHASE_ORDER_BODY)
        terms = ("Terms: net 30. " * 400).encode()  # a regular-sector stream, beside mini streams
        content = make_msg([("po.docx", DOCX, word), ("quote.xlsx", XLSX, make_xlsx()),
                            ("terms.txt", "text/plain", terms),
                            ("fwd.msg", "application/vnd.ms-outlook", b"nested")])
        document = self.extract(content, "application/octet-stream", backend=None,
                                filename="confirmation.msg")
        self.assertEqual((document.status, document.method, document.warnings),
                         ("extracted", "outlook-msg", ()))
        self.assertEqual(document.text,
            "Subject: PO PPO-200 confirmation\nFrom: Acme Sales <sales@acme.example>\n"
            "To: Purchasing\nDate: Wed, 30 Sep 2026 10:00:00 +0000\nPetra PO: PPO-200\r\n"
            "See the attached quote.")
        self.assertEqual((document.conversion, document.sheets), (None, ()))
        # Each attachment is a member, read on its own bytes with its own provenance.
        files = {f["name"]: f for f in use_case_parts.members(content, "confirmation.msg")["files"]}
        self.assertEqual(set(files), {"po.docx", "quote.xlsx", "terms.txt", "fwd.msg"})
        self.assertEqual(files["terms.txt"]["data"], terms)
        po = self.extract(files["po.docx"]["data"], DOCX, backend=None)
        self.assertEqual((po.text, po.conversion["source_sha256"]),
                         (PURCHASE_ORDER_MD, sha256(word).hexdigest()))
        self.assertTrue(stored_provenance_valid(po.evidence_attributes(URL)))
        self.assertEqual([(c["field"], c.get("converted", False)) for c in po.claims
                          if c["field"] == "petra_po"], [("petra_po", True)])
        quote = self.extract(files["quote.xlsx"]["data"], XLSX, backend=None)
        sheet, = quote.sheets
        self.assertEqual(sheet["sheet"], "Quote")
        self.assertEqual(next(c for c in quote.claims if c["field"] == "quote_number")["sheet"],
                         "Quote")

    def test_outlook_msg_with_only_an_rtf_body_is_read_as_text(self):
        from operating.tests.test_use_case_intake import rtf_compressed
        rtf = rtf_compressed(b"{\\rtf1\\ansi\\ansicpg1252{\\fonttbl{\\f0 Arial;}}"
                             b"Petra PO: PPO-300\\par Ship date moved\\par}")
        document = self.extract(make_msg([], rtf=rtf), "application/vnd.ms-outlook",
                                backend=None)
        self.assertEqual((document.status, document.warnings), ("extracted", ()))
        self.assertTrue(document.text.endswith("+0000\nPetra PO: PPO-300\nShip date moved"))
        self.assertEqual([c["value"] for c in document.claims if c["field"] == "petra_po"],
                         ["PPO-300"])
        damaged = self.extract(make_msg([], rtf=rtf[:-3]), "application/vnd.ms-outlook",
                               backend=None)
        self.assertEqual((damaged.status, damaged.warnings),
                         ("extracted", ("msg_rtf_body_unreadable",)))
        self.assertNotIn("PPO-300", damaged.text)

    def test_a_fat_sector_listed_twice_is_refused_not_multiplied(self):
        # 109 DIFAT slots naming one sector once grew the FAT 109-fold (a 1 MB file
        # could allocate hundreds of MB); the one reader refuses the file instead.
        content = bytearray(make_msg([]))
        struct.pack_into("<I", content, 0x2C, 109)
        struct.pack_into("<109I", content, 0x4C, *[0] * 109)
        document = self.extract(bytes(content), "application/vnd.ms-outlook", backend=None)
        self.assertEqual((document.status, document.warnings),
                         ("unsupported", ("msg_invalid_compound_file",)))

    def test_the_repo_has_one_compound_file_reader(self):
        root = Path(__file__).resolve().parents[2]
        readers, magics = [], []
        for path in sorted(root.rglob("*.py")):
            relative = path.relative_to(root)
            if any(part.startswith(".") or part in ("tests", "EVIDENCE")
                   for part in relative.parts):
                continue  # test fixtures build compound files byte by byte
            source = path.read_text(encoding="utf-8", errors="replace")
            readers += [f"{relative}:{name}" for name in
                        re.findall(r"(?m)^\s*class (\w*cfb\w*)\b", source, re.I)]
            if re.search(r"\\xd0\\xcf\\x11\\xe0|d0\s*cf\s*11\s*e0", source, re.I):
                magics.append(str(relative))
        self.assertEqual((readers, magics),
                         (["operating/msg_reader.py:Cfb"], ["operating/msg_reader.py"]))


class InstalledStateTest(unittest.TestCase):
    def setUp(self):
        doc_convert._import_backend.cache_clear()
        self.addCleanup(doc_convert._import_backend.cache_clear)

    def test_backend_reports_the_installed_state_honestly(self):
        # Found-but-broken (a native dependency that panics on import) is a real host
        # state, so the expectation comes from importing in a subprocess the test
        # cannot be killed by, not from find_spec alone.
        if importlib.util.find_spec("markitdown") is None:
            expected = False
        else:
            probe = subprocess.run([sys.executable, "-I", "-c", "from markitdown import MarkItDown"],
                                   capture_output=True, timeout=240)
            expected = probe.returncode == 0
        self.assertEqual(doc_convert.available(), expected)

    def test_import_outcomes_map_to_available(self):
        class Panic(BaseException):  # like pyo3's PanicException
            pass

        class Fake:
            def __init__(self, enable_plugins=False):
                pass

        good = types.ModuleType("markitdown")
        good.MarkItDown = Fake
        with patch.dict(sys.modules, {"markitdown": good}):
            self.assertTrue(doc_convert.available())
        real_import = builtins.__import__

        def importing_raises(exc):
            def fake(name, *a, **k):
                if name == "markitdown":
                    raise exc
                return real_import(name, *a, **k)
            return fake

        with patch.dict(sys.modules, {"markitdown": None}):  # not importable
            doc_convert._import_backend.cache_clear()
            self.assertFalse(doc_convert.available())
        for exc in (ImportError(), Panic()):
            doc_convert._import_backend.cache_clear()
            with patch("builtins.__import__", importing_raises(exc)):
                self.assertFalse(doc_convert.available())
        for exc in (KeyboardInterrupt(), SystemExit()):  # never swallowed
            doc_convert._import_backend.cache_clear()
            with patch("builtins.__import__", importing_raises(exc)):
                with self.assertRaises(type(exc)):
                    doc_convert.available()


if __name__ == "__main__":
    unittest.main()
